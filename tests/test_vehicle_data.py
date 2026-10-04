import numpy as np
import openpyxl

import flight_acoustics as fa


def _write_reflist_vehicle(directory, rows):
    """A vehicle.cfg with an [Option] reflist, and the reflist workbook.

    ``rows`` are (run, speed_kt, mu, cw, mtip); a run of None leaves column B
    blank on that row.
    """
    (directory / 'vehicle.cfg').write_text(
        '[Main Rotor]\nblades = 4\n[Tail Rotor]\nblades = 2\n'
        '[Option]\nfbar = 0.01\nreflist = runs.xlsx\n')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['B1'] = 'Run'
    for i, (run, speed, mu, cw, mtip) in enumerate(rows, start=2):
        if run is not None:
            ws['B{}'.format(i)] = run
        ws['W{}'.format(i)] = speed
        ws['AH{}'.format(i)] = mu
        ws['AI{}'.format(i)] = cw
        ws['AJ{}'.format(i)] = mtip
    wb.save(directory / 'runs.xlsx')


def test_reflist_rows_stay_aligned_past_a_blank_run_cell(tmp_path):
    _write_reflist_vehicle(tmp_path, [
        (101, 60.0, 0.10, 0.006, 0.60),
        (None, 0.0, 0.0, 0.0, 0.0),
        (102, 80.0, 0.14, 0.007, 0.61),
        (103, 100.0, 0.18, 0.008, 0.62),
    ])
    out = fa.read_vehicle_data(str(tmp_path), runs=np.array([103, 101]), speeds=None,
                               flight_path_angles=np.array([3.0, -6.0]))
    advance_ratios, weight_coefficients, hover_tip_mach_numbers, alphas = out[8], out[9], out[10], out[11]
    np.testing.assert_array_equal(advance_ratios, [0.18, 0.10])
    np.testing.assert_array_equal(weight_coefficients, [0.008, 0.006])
    np.testing.assert_array_equal(hover_tip_mach_numbers, [0.62, 0.60])
    np.testing.assert_array_equal(out[13], [100.0, 60.0])


def test_reflist_gives_angles_of_attack(tmp_path):
    _write_reflist_vehicle(tmp_path, [(101, 60.0, 0.10, 0.006, 0.60)])
    out = fa.read_vehicle_data(str(tmp_path), runs=np.array([101]), speeds=None,
                               flight_path_angles=np.array([-6.0]))
    alphas = out[11]
    expected = -np.degrees(0.5 * 0.01 * 0.10 ** 2 / 0.006) + 6.0
    np.testing.assert_allclose(alphas, [expected])


def test_vehicle_weight_uses_standard_gravity(tmp_path):
    (tmp_path / 'vehicle.cfg').write_text('[Vehicle]\nweight = 1000\n')
    assert fa.read_vehicle_weight_newtons(str(tmp_path)) == 9806.65
